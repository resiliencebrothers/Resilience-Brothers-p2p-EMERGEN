"""iter220 — Exportar Inventario: CSV (control / movimientos con filtros) y
Excel completo (Control Inventario + Movimientos del período + Dashboard).
"""
import io
import os
import time

import requests
from openpyxl import load_workbook
from pymongo import MongoClient

from tests.conftest import BASE_URL, ADMIN_TOKEN, VIP_TOKEN

API = f"{BASE_URL}/api"
MARK = "ITER220TEST"


def _h(tok=None):
    h = {"Content-Type": "application/json"}
    if tok:
        h["Authorization"] = f"Bearer {tok}"
    return h


def _db():
    return MongoClient(os.environ["MONGO_URL"])[os.environ["DB_NAME"]]


def _cleanup():
    db = _db()
    ids = [p["id"] for p in db.products.find({"category": "iter220test"}, {"id": 1})]
    db.products.delete_many({"category": "iter220test"})
    if ids:
        db.inventory_movements.delete_many({"product_id": {"$in": ids}})
        db.notifications.delete_many({"data.product_id": {"$in": ids}})


def setup_module(module):
    _cleanup()


def teardown_module(module):
    _cleanup()


def _create_product_with_entrada():
    name = f"{MARK}_Prod_{int(time.time()*1000)}"
    r = requests.post(f"{API}/admin/products", headers=_h(ADMIN_TOKEN), json={
        "name": name, "description": "iter220", "image_url": "",
        "price_usd": 15.0, "cost_usd": 9.0, "stock": 10,
        "category": "iter220test", "is_active": True})
    assert r.status_code == 200, r.text
    p = r.json()
    r = requests.post(f"{API}/admin/inventory/movements", headers=_h(ADMIN_TOKEN),
                      json={"product_id": p["id"], "type": "entrada",
                            "quantity": 5, "note": MARK})
    assert r.status_code == 200, r.text
    return p


class TestCsvExport:
    def test_control_csv_contains_products(self):
        p = _create_product_with_entrada()
        r = requests.get(f"{API}/admin/inventory/export.csv",
                         params={"dataset": "control"}, headers=_h(ADMIN_TOKEN))
        assert r.status_code == 200, r.text
        assert "text/csv" in r.headers["content-type"]
        assert "attachment" in r.headers["content-disposition"]
        body = r.content.decode("utf-8-sig")
        assert "Producto" in body and "Existencia" in body
        assert p["name"] in body
        assert "Maicena" in body  # producto del seed Excel

    def test_movements_csv_with_filters(self):
        p = _create_product_with_entrada()
        r = requests.get(f"{API}/admin/inventory/export.csv",
                         params={"dataset": "movements",
                                 "product_id": p["id"]},
                         headers=_h(ADMIN_TOKEN))
        assert r.status_code == 200, r.text
        body = r.content.decode("utf-8-sig")
        lines = [line for line in body.strip().splitlines() if line]
        assert "Ganancia" in lines[0] and "Usuario" in lines[0]
        # iter226 — el alta con stock inicial también queda como entrada:
        # header + entrada 'alta' + la entrada manual de este producto.
        assert len(lines) == 3
        assert p["name"] in lines[1] and "Entrada" in lines[1]
        assert "admin.test@resilience.com" in lines[1]

    def test_invalid_dataset_rejected(self):
        r = requests.get(f"{API}/admin/inventory/export.csv",
                         params={"dataset": "otro"}, headers=_h(ADMIN_TOKEN))
        assert r.status_code == 400

    def test_client_cannot_export(self):
        r = requests.get(f"{API}/admin/inventory/export.csv",
                         params={"dataset": "control"}, headers=_h(VIP_TOKEN))
        assert r.status_code == 403
        r = requests.get(f"{API}/admin/inventory/export.xlsx",
                         headers=_h(VIP_TOKEN))
        assert r.status_code == 403


class TestXlsxExport:
    def test_full_workbook(self):
        p = _create_product_with_entrada()
        today = time.strftime("%Y-%m-%d")
        r = requests.get(f"{API}/admin/inventory/export.xlsx",
                         params={"start": today, "end": today},
                         headers=_h(ADMIN_TOKEN))
        assert r.status_code == 200, r.text
        assert "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Control Inventario", "Movimientos", "Dashboard"]
        control_names = {row[0].value for row in wb["Control Inventario"].iter_rows(min_row=2)}
        assert p["name"] in control_names
        assert "Maicena" in control_names
        mov_products = {row[2].value for row in wb["Movimientos"].iter_rows(min_row=2)}
        assert p["name"] in mov_products
        dash = {row[0].value: row[1].value for row in wb["Dashboard"].iter_rows()}
        assert "Ganancia realizada" in dash
        assert "Rentabilidad s/ costo %" in dash
        assert "Comisión ventas VIP" in dash

    def test_invalid_range_rejected(self):
        r = requests.get(f"{API}/admin/inventory/export.xlsx",
                         params={"start": "2026-09-01", "end": "2026-08-01"},
                         headers=_h(ADMIN_TOKEN))
        assert r.status_code == 400
