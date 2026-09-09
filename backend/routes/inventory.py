"""Control de inventario de la tienda física — iter217.

Endpoints (todos staff con permiso 'products'):
- GET  /admin/inventory/control              tabla en tiempo real por producto
- GET  /admin/inventory/movements            registro de entradas y salidas
- POST /admin/inventory/movements            registrar movimiento manual
- GET  /admin/inventory/dashboard            KPIs del período
"""
import logging
from typing import Optional, Literal, Any

from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel, Field

from db_client import db
from auth_utils import require_permission, now_utc, iso
from audit_log import log_action
from services.inventory import (record_movement, record_price_change,
                                _day_bounds, today_havana,
                                build_control_rows, build_dashboard,
                                build_rotation)
from services.proof_upload import maybe_upload_proof

logger = logging.getLogger(__name__)
router = APIRouter(tags=["Inventory"])

_TYPE_LABELS = {"entrada": "Entrada", "venta": "Venta",
                "ajuste_pos": "Ajuste +", "ajuste_neg": "Ajuste −"}


class MovementCreate(BaseModel):
    # iter219 — los ajustes manuales se retiraron a pedido del operador:
    # solo Entrada y Venta (los reversos internos siguen usando ajuste_pos).
    product_id: str
    type: Literal["entrada", "venta"]
    quantity: int = Field(..., gt=0, le=1_000_000)
    unit_price: Optional[float] = Field(None, ge=0)
    unit_cost: Optional[float] = Field(None, ge=0)
    # iter219 — en una Entrada permite actualizar el precio de venta del
    # producto (ficha de costo del Excel). Queda auditado como tipo 'precio'.
    sale_price: Optional[float] = Field(None, ge=0)
    note: str = Field("", max_length=300)
    # iter219 — foto opcional del producto en oferta (base64 data URL → R2).
    photo_url: str = ""


@router.get("/admin/inventory/control")
async def inventory_control(request: Request) -> Any:
    await require_permission(request, "products")
    return await build_control_rows()


@router.get("/admin/inventory/movements")
async def list_movements(request: Request, product_id: Optional[str] = None,
                         type: Optional[str] = None,
                         start: Optional[str] = None,
                         end: Optional[str] = None,
                         limit: int = 200) -> Any:
    await require_permission(request, "products")
    q: dict = {}
    if product_id:
        q["product_id"] = product_id
    if type:
        q["type"] = type
    created: dict = {}
    if start:
        created["$gte"] = _day_bounds(start)[0]
    if end:
        created["$lt"] = _day_bounds(end)[1]
    if created:
        q["created_at"] = created
    limit = max(1, min(int(limit), 500))
    return await db.inventory_movements.find(q, {"_id": 0}) \
        .sort("created_at", -1).to_list(limit)


@router.post("/admin/inventory/movements")
async def create_movement(payload: MovementCreate, request: Request) -> Any:
    actor = await require_permission(request, "products")
    product = await db.products.find_one({"id": payload.product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    if product.get("owner_id"):
        raise HTTPException(
            status_code=400,
            detail="Los productos de vendedores VIP no entran al inventario de la empresa")
    photo = maybe_upload_proof(payload.photo_url, "inventory") or "" if payload.photo_url else ""
    doc = await record_movement(
        product=product, mtype=payload.type, quantity=payload.quantity,
        unit_price=payload.unit_price, unit_cost=payload.unit_cost,
        note=payload.note, source="manual", actor=actor, photo_url=photo)
    # iter219 — una Entrada puede actualizar la ficha del producto (precio de
    # venta / costo unitario). Cada cambio queda auditado (tipo 'precio').
    if payload.type == "entrada":
        updates = {}
        if payload.sale_price is not None and \
                float(payload.sale_price) != float(product.get("price_usd") or 0):
            updates["price_usd"] = round(float(payload.sale_price), 2)
            await record_price_change(
                product=product, field_label="Precio venta",
                old=float(product.get("price_usd") or 0),
                new=updates["price_usd"], actor=actor)
        if payload.unit_cost is not None and \
                float(payload.unit_cost) != float(product.get("cost_usd") or 0):
            updates["cost_usd"] = round(float(payload.unit_cost), 2)
            await record_price_change(
                product=product, field_label="Costo unitario",
                old=float(product.get("cost_usd") or 0),
                new=updates["cost_usd"], actor=actor)
        if updates:
            await db.products.update_one({"id": product["id"]}, {"$set": updates})
    await log_action(
        db, actor, f"inventory.{payload.type}", "inventory_movement", doc["id"],
        summary=(f"{_TYPE_LABELS[payload.type]} {payload.quantity}× "
                 f"{product.get('name', '')}"),
        details={"product_id": product["id"], "quantity": payload.quantity,
                 "total": doc["total"], "note": payload.note})
    try:
        from services.live_bus import publish
        await publish("products_changed", {"product_id": product["id"]})
    except Exception as e:  # noqa: BLE001
        logger.error(f"products_changed publish failed: {e}")
    return doc


class BarcodeAssign(BaseModel):
    product_id: str
    barcode: str = Field(..., min_length=3, max_length=64)


class BarcodeGenerate(BaseModel):
    product_ids: list = Field(..., min_length=1, max_length=500)


def _ean13_check_digit(d12: str) -> str:
    s = sum(int(c) * (3 if i % 2 else 1) for i, c in enumerate(d12))
    return str((10 - s % 10) % 10)


async def _generate_internal_ean13() -> str:
    """iter228 — EAN-13 interno (prefijo 20x, reservado para uso en tienda)."""
    import secrets
    for _ in range(25):
        body = "200" + "".join(secrets.choice("0123456789") for _ in range(9))
        code = body + _ean13_check_digit(body)
        if not await db.products.find_one({"barcode": code}, {"_id": 1}):
            return code
    raise HTTPException(status_code=500,
                        detail="No se pudo generar un código único")


_COMPANY_Q = {"$or": [{"owner_id": {"$in": [None, ""]}},
                      {"owner_id": {"$exists": False}}]}


@router.get("/admin/inventory/barcode/{code}")
async def lookup_barcode(code: str, request: Request) -> Any:
    """iter227 — escáner: busca el producto vinculado a un código de barras."""
    await require_permission(request, "products")
    code = code.strip()
    p = await db.products.find_one({"barcode": code, **_COMPANY_Q}, {"_id": 0})
    if not p:
        raise HTTPException(status_code=404,
                            detail="Código no vinculado a ningún producto")
    return {"product_id": p["id"], "name": p.get("name", ""),
            "stock": int(p.get("stock") or 0),
            "price_usd": float(p.get("price_usd") or 0),
            "cost_usd": float(p.get("cost_usd") or 0),
            "image_url": p.get("image_url", ""),
            "is_active": bool(p.get("is_active", True)),
            "barcode": code}


@router.post("/admin/inventory/barcode/assign")
async def assign_barcode(payload: BarcodeAssign, request: Request) -> Any:
    """iter227 — vincula un código de barras a un producto de la empresa."""
    actor = await require_permission(request, "products")
    code = payload.barcode.strip()
    if not code:
        raise HTTPException(status_code=400, detail="Código de barras vacío")
    product = await db.products.find_one({"id": payload.product_id}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Producto no encontrado")
    if product.get("owner_id"):
        raise HTTPException(
            status_code=400,
            detail="Los productos de vendedores VIP no entran al inventario de la empresa")
    other = await db.products.find_one(
        {"barcode": code, "id": {"$ne": product["id"]}, **_COMPANY_Q}, {"_id": 0})
    if other:
        raise HTTPException(status_code=409, detail=(
            f"Este código ya está vinculado a «{other.get('name', '')}»"))
    await db.products.update_one({"id": product["id"]},
                                 {"$set": {"barcode": code}})
    await log_action(
        db, actor, "inventory.barcode_assign", "product", product["id"],
        summary=f"Código {code} → {product.get('name', '')}",
        details={"barcode": code})
    return {"ok": True, "product_id": product["id"], "barcode": code,
            "name": product.get("name", "")}


@router.post("/admin/inventory/barcode/generate")
async def generate_barcodes(payload: BarcodeGenerate, request: Request) -> Any:
    """iter228 — genera códigos internos para productos sin código de barras."""
    actor = await require_permission(request, "products")
    out = []
    for pid in payload.product_ids:
        product = await db.products.find_one({"id": str(pid)}, {"_id": 0})
        if not product or product.get("owner_id"):
            continue
        if product.get("barcode"):
            out.append({"product_id": product["id"],
                        "barcode": product["barcode"], "generated": False})
            continue
        code = await _generate_internal_ean13()
        await db.products.update_one({"id": product["id"]},
                                     {"$set": {"barcode": code}})
        await log_action(
            db, actor, "inventory.barcode_generate", "product", product["id"],
            summary=f"Código interno {code} → {product.get('name', '')}",
            details={"barcode": code})
        out.append({"product_id": product["id"], "barcode": code,
                    "generated": True})
    return out


@router.get("/admin/inventory/labels.pdf")
async def labels_pdf(request: Request, product_ids: str,
                     copies: int = 1) -> Any:
    """iter228 — PDF A4 (rejilla 3×8) de etiquetas con código de barras."""
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from services.label_pdf import build_labels_pdf

    await require_permission(request, "products")
    ids = [x.strip() for x in product_ids.split(",") if x.strip()]
    if not ids:
        raise HTTPException(status_code=400, detail="Sin productos seleccionados")
    copies = max(1, min(int(copies), 50))
    prods = await db.products.find(
        {"id": {"$in": ids}, **_COMPANY_Q}, {"_id": 0}).to_list(500)
    items = sorted([p for p in prods if p.get("barcode")],
                   key=lambda p: (p.get("name") or "").lower())
    if not items:
        raise HTTPException(
            status_code=400,
            detail="Ninguno de los productos seleccionados tiene código de barras")
    pdf = build_labels_pdf(items, copies)
    ts = iso(now_utc())[:10]
    return StreamingResponse(
        BytesIO(pdf), media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="etiquetas_{ts}.pdf"'})


@router.get("/admin/inventory/daily-close")
async def daily_close(request: Request, date: Optional[str] = None) -> Any:
    """iter230 — resumen de caja del día para cuadrar al cerrar la tienda."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    from services.inventory import build_daily_close

    await require_permission(request, "products")
    if not date:
        date = datetime.now(ZoneInfo("America/Havana")).strftime("%Y-%m-%d")
    try:
        datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400,
                            detail="Fecha inválida (formato YYYY-MM-DD)")
    return await build_daily_close(date)


@router.get("/admin/inventory/daily-close.pdf")
async def daily_close_pdf(request: Request, date: Optional[str] = None) -> Any:
    """iter231 — cierre del día en PDF para archivar o compartir con socios."""
    from datetime import datetime
    from io import BytesIO
    from zoneinfo import ZoneInfo
    from fastapi.responses import StreamingResponse
    from services.inventory import build_daily_close
    from store_close_pdf import build_store_close_pdf

    await require_permission(request, "products")
    if not date:
        date = datetime.now(ZoneInfo("America/Havana")).strftime("%Y-%m-%d")
    try:
        datetime.strptime(date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400,
                            detail="Fecha inválida (formato YYYY-MM-DD)")
    pdf = build_store_close_pdf(await build_daily_close(date))
    return StreamingResponse(
        BytesIO(pdf), media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="cierre_tienda_{date}.pdf"'})


@router.get("/admin/inventory/dashboard")
async def inventory_dashboard(request: Request, start: Optional[str] = None,
                              end: Optional[str] = None,
                              product_id: Optional[str] = None,
                              product_ids: Optional[str] = None) -> Any:
    """iter224 — `product_ids` (separados por coma) permite analizar varias
    mercancías a la vez; `product_id` se mantiene por compatibilidad."""
    await require_permission(request, "products")
    today = today_havana()
    start = (start or today)[:10]
    end = (end or today)[:10]
    if start > end:
        raise HTTPException(status_code=400, detail="Rango de fechas inválido")
    ids = [x.strip() for x in (product_ids or "").split(",") if x.strip()]
    if not ids and product_id:
        ids = [product_id]
    return await build_dashboard(start, end, product_ids=ids or None)


@router.get("/admin/inventory/rotation")
async def inventory_rotation(request: Request, start: Optional[str] = None,
                             end: Optional[str] = None,
                             product_ids: Optional[str] = None) -> Any:
    """iter225 — rotación por mercancía: cuánto tarda cada producto en
    venderse (ritmo/día, días para agotarse) y en reponerse (ciclo entre
    entradas)."""
    await require_permission(request, "products")
    today = today_havana()
    start = (start or today)[:10]
    end = (end or today)[:10]
    if start > end:
        raise HTTPException(status_code=400, detail="Rango de fechas inválido")
    ids = [x.strip() for x in (product_ids or "").split(",") if x.strip()]
    return await build_rotation(start, end, product_ids=ids or None)


async def _movement_rows(product_id: Optional[str], type: Optional[str],
                         start: Optional[str], end: Optional[str]) -> list:
    q: dict = {}
    if product_id:
        q["product_id"] = product_id
    if type:
        q["type"] = type
    created: dict = {}
    if start:
        created["$gte"] = _day_bounds(start)[0]
    if end:
        created["$lt"] = _day_bounds(end)[1]
    if created:
        q["created_at"] = created
    return await db.inventory_movements.find(q, {"_id": 0}) \
        .sort("created_at", -1).to_list(20000)


@router.get("/admin/inventory/export.csv")
async def export_inventory_csv(request: Request, dataset: str = "movements",
                               product_id: Optional[str] = None,
                               type: Optional[str] = None,
                               start: Optional[str] = None,
                               end: Optional[str] = None) -> Any:
    """iter220 — CSV contable del control de inventario o del registro de
    movimientos (con filtros de producto/tipo/fechas)."""
    import csv
    import io
    from io import BytesIO
    from fastapi.responses import StreamingResponse

    await require_permission(request, "products")
    text_buf = io.StringIO()
    writer = csv.writer(text_buf, quoting=csv.QUOTE_ALL)
    if dataset == "control":
        writer.writerow(["Producto", "Categoría", "Activo", "Existencia",
                         "Precio venta", "Costo unitario", "Entradas",
                         "Ventas", "Valor inventario", "Vendido hoy",
                         "Ingresos hoy", "Estado"])
        for r in await build_control_rows():
            writer.writerow([
                r["name"], r["category"], "sí" if r["is_active"] else "no",
                r["stock"], r["price_usd"], r["cost_usd"], r["entradas"],
                r["ventas"], r["inventory_value"], r["sold_today"],
                r["revenue_today"], r["estado"].upper()])
        prefix = "inventario_control"
    elif dataset == "movements":
        writer.writerow(["Fecha", "Tipo", "Producto", "Cantidad",
                         "Precio unitario", "Costo unitario", "Total",
                         "Costo de venta", "Ganancia", "Fuente", "Usuario",
                         "Nota", "Foto"])
        for m in await _movement_rows(product_id, type, start, end):
            writer.writerow([
                m["created_at"], _TYPE_LABELS.get(m["type"], m["type"]),
                m["product_name"], m["quantity"], m["unit_price"],
                m["unit_cost"], m["total"], m["cost_of_sale"], m["profit"],
                m.get("source", ""), m.get("actor_email", ""),
                m.get("note", ""), m.get("photo_url", "")])
        prefix = "inventario_movimientos"
    else:
        raise HTTPException(status_code=400, detail="dataset inválido (control|movements)")

    buf = BytesIO()
    buf.write(text_buf.getvalue().encode("utf-8-sig"))
    buf.seek(0)
    ts = iso(now_utc())[:16].replace("-", "").replace(":", "").replace("T", "_")
    return StreamingResponse(
        buf, media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="{prefix}_{ts}.csv"'})


@router.get("/admin/inventory/export.xlsx")
async def export_inventory_xlsx(request: Request,
                                start: Optional[str] = None,
                                end: Optional[str] = None) -> Any:
    """iter220 — Excel contable completo: Control Inventario + Movimientos
    del período + Dashboard (por defecto: mes en curso)."""
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    await require_permission(request, "products")
    today = today_havana()
    start = (start or f"{today[:7]}-01")[:10]
    end = (end or today)[:10]
    if start > end:
        raise HTTPException(status_code=400, detail="Rango de fechas inválido")

    wb = Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Control Inventario"
    ws.append(["Producto", "Categoría", "Activo", "Existencia", "Precio venta",
               "Costo unitario", "Entradas", "Ventas", "Valor inventario",
               "Vendido hoy", "Ingresos hoy", "Estado"])
    for cell in ws[1]:
        cell.font = bold
    for r in await build_control_rows():
        ws.append([r["name"], r["category"], "sí" if r["is_active"] else "no",
                   r["stock"], r["price_usd"], r["cost_usd"], r["entradas"],
                   r["ventas"], r["inventory_value"], r["sold_today"],
                   r["revenue_today"], r["estado"].upper()])
    ws.column_dimensions["A"].width = 32

    ws2 = wb.create_sheet("Movimientos")
    ws2.append(["Fecha", "Tipo", "Producto", "Cantidad", "Precio unitario",
                "Costo unitario", "Total", "Costo de venta", "Ganancia",
                "Fuente", "Usuario", "Nota"])
    for cell in ws2[1]:
        cell.font = bold
    for m in await _movement_rows(None, None, start, end):
        ws2.append([m["created_at"], _TYPE_LABELS.get(m["type"], m["type"]),
                    m["product_name"], m["quantity"], m["unit_price"],
                    m["unit_cost"], m["total"], m["cost_of_sale"], m["profit"],
                    m.get("source", ""), m.get("actor_email", ""),
                    m.get("note", "")])
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["C"].width = 32

    ws3 = wb.create_sheet("Dashboard")
    d = await build_dashboard(start, end)
    kpis = [
        ("Período", f"{start} → {end}"),
        ("Unidades vendidas", d["units_sold"]),
        ("Ingresos por ventas", d["sales_revenue"]),
        ("Costo mercancía vendida", d["cogs"]),
        ("Ganancia realizada", d["profit"]),
        ("Margen realizado %", d["margin_pct"]),
        ("Rentabilidad s/ costo %", d["profitability_pct"]),
        ("Salidas por compras", d["purchases_out"]),
        ("Flujo neto de caja", d["net_cash_flow"]),
        ("Valor inventario actual", d["inventory_value"]),
        ("Unidades en existencia", d["units_in_stock"]),
        ("Comisión ventas VIP", d["vip_commission_earned"] or 0),
    ]
    for label, value in kpis:
        ws3.append([label, value])
    for row in ws3.iter_rows(min_col=1, max_col=1):
        row[0].font = bold
    ws3.column_dimensions["A"].width = 30

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="inventario_{start}_a_{end}.xlsx"'})
