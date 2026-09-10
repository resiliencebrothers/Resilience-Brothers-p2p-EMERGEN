"""iter233 — Caja de Efectivo ("Caja En Vivo" integrada a la plataforma).

Control de efectivo físico CUP/USD sin Excel, para TODOS los usuarios:
- Cajas PERSONALES (privadas de cada cliente/staff) y cajas de EMPRESA
  (compartidas, solo admin/employee).
- Cada caja maneja dos fondos: CUP y USD.
- Saldo inicial con desglose de billetes, movimientos entrada/salida
  (concepto, responsable, importe, desglose opcional), balance y neto del
  mes, control físico por denominación, arqueo (contado vs sistema),
  reporte mensual por días y respaldo Excel.

Nota: es una herramienta de control físico independiente del Fondo de
Empresa contable (admin_company_funds).
"""
import uuid
from datetime import datetime
from typing import Any, Dict, Literal, Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, HTTPException, Request
from pydantic import BaseModel, Field

from db_client import db
from auth_utils import require_user, iso, now_utc
# núcleo compartido — los servicios no importan rutas (review iter270)
from services.cash_box_core import DENOMS, FUNDS, fund_balance as _fund_balance

router = APIRouter(tags=["CashBox"])

_TZ = ZoneInfo("America/Havana")


class BoxCreate(BaseModel):
    name: str = Field(..., min_length=2, max_length=60)
    scope: Literal["personal", "empresa"] = "personal"


class BoxUpdate(BaseModel):
    name: Optional[str] = Field(default=None, min_length=2, max_length=60)
    is_active: Optional[bool] = None


class InitialSet(BaseModel):
    fund: Literal["CUP", "USD"]
    amount: float = Field(..., ge=0)
    denominations: Optional[Dict[str, int]] = None


class MovementCreate(BaseModel):
    fund: Literal["CUP", "USD"]
    type: Literal["entrada", "salida"]
    amount: float = Field(..., gt=0)
    concept: str = Field(..., min_length=1, max_length=200)
    responsible: str = Field(default="", max_length=80)
    denominations: Optional[Dict[str, int]] = None


class MovementUpdate(BaseModel):
    amount: Optional[float] = Field(default=None, gt=0)
    concept: Optional[str] = Field(default=None, min_length=1, max_length=200)
    responsible: Optional[str] = Field(default=None, max_length=80)
    denominations: Optional[Dict[str, int]] = None


class ArqueoCreate(BaseModel):
    fund: Literal["CUP", "USD"]
    counted: Dict[str, int]
    note: str = Field(default="", max_length=200)


def _is_staff(user: dict) -> bool:
    return user.get("role") in ("admin", "employee")


# H01/V02 — movimientos espejados desde el Fondo de Empresa (origen contable)
_LINK_FIELDS = ("source_adjustment_id", "source_withdrawal_id",
                "source_client_withdrawal_id", "source_transfer_id")


def _ledger_linked(mov: dict) -> bool:
    return any(mov.get(k) for k in _LINK_FIELDS)


async def _bump_fund_rev(box_id: str, fund: str) -> None:
    """V04 — revisión del fondo: crear/editar/eliminar movimientos invalida
    el arqueo de cierre vigente (el conteo anterior queda como histórico)."""
    await db.cash_boxes.update_one(
        {"id": box_id}, {"$inc": {f"fund_revs.{fund}": 1}})


def _clean_denoms(fund: str, raw: Optional[Dict[str, int]],
                  amount: Optional[float] = None) -> Optional[Dict[str, int]]:
    """Valida el desglose: denominaciones válidas del fondo y, si se pasa
    `amount`, que el total del desglose coincida con el importe."""
    if not raw:
        return None
    valid = DENOMS[fund]
    clean: Dict[str, int] = {}
    total = 0.0
    for k, v in raw.items():
        try:
            denom, qty = int(float(k)), int(v)
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Desglose inválido")
        if denom not in valid:
            raise HTTPException(
                status_code=400,
                detail=f"Denominación {denom} no válida para {fund}")
        if qty < 0:
            raise HTTPException(status_code=400,
                                detail="Las cantidades no pueden ser negativas")
        if qty:
            # H05 — claves equivalentes (100 / "100.0") se CONSOLIDAN: nunca
            # se pierde una cantidad ya sumada al total validado.
            clean[str(denom)] = clean.get(str(denom), 0) + qty
            total += denom * qty
    if amount is not None and clean and round(total, 2) != round(float(amount), 2):
        raise HTTPException(
            status_code=400,
            detail=(f"El desglose suma {total:g} y el importe es "
                    f"{float(amount):g}. Deben coincidir."))
    return clean or None


async def _get_box_checked(box_id: str, user: dict) -> dict:
    box = await db.cash_boxes.find_one({"id": box_id}, {"_id": 0})
    if not box:
        raise HTTPException(status_code=404, detail="Caja no encontrada")
    if box["scope"] == "personal":
        if box["owner_id"] != user["user_id"]:
            raise HTTPException(status_code=403, detail="No es tu caja")
    elif not _is_staff(user):
        raise HTTPException(status_code=403,
                            detail="Solo el staff maneja las cajas de empresa")
    return box


def _month_range_utc(month: str) -> tuple:
    day = datetime.strptime(month + "-01", "%Y-%m-%d").replace(tzinfo=_TZ)
    if day.month == 12:
        nxt = day.replace(year=day.year + 1, month=1)
    else:
        nxt = day.replace(month=day.month + 1)
    from datetime import timezone as _tzu
    return (day.astimezone(_tzu.utc).isoformat(),
            nxt.astimezone(_tzu.utc).isoformat())


# ---------------------------------------------------------------- cajas

@router.get("/cashbox/boxes")
async def list_boxes(request: Request) -> Any:
    user = await require_user(request)
    q: Dict[str, Any] = {"$or": [
        {"scope": "personal", "owner_id": user["user_id"]}]}
    if _is_staff(user):
        q["$or"].append({"scope": "empresa"})
    boxes = await db.cash_boxes.find(q, {"_id": 0}).sort("created_at", 1).to_list(200)
    for b in boxes:
        b["balances"] = {f: await _fund_balance(b, f) for f in FUNDS}
    return boxes


@router.post("/cashbox/boxes")
async def create_box(payload: BoxCreate, request: Request) -> Any:
    user = await require_user(request)
    if payload.scope == "empresa" and not _is_staff(user):
        raise HTTPException(status_code=403,
                            detail="Solo el staff crea cajas de empresa")
    doc = {
        "id": f"cbox_{uuid.uuid4().hex[:12]}",
        "name": payload.name.strip(),
        "scope": payload.scope,
        "owner_id": user["user_id"] if payload.scope == "personal" else "",
        "created_by_id": user["user_id"],
        "created_by_name": user.get("name") or user.get("email") or "",
        "created_at": iso(now_utc()),
        "is_active": True,
        "initial": {},
    }
    await db.cash_boxes.insert_one({**doc})
    doc["balances"] = {f: 0.0 for f in FUNDS}
    return doc


@router.put("/cashbox/boxes/{box_id}")
async def update_box(box_id: str, payload: BoxUpdate, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    upd: dict = {}
    if payload.name is not None:
        upd["name"] = payload.name.strip()
    if payload.is_active is not None:
        upd["is_active"] = payload.is_active
    if upd:
        await db.cash_boxes.update_one({"id": box["id"]}, {"$set": upd})
    return await db.cash_boxes.find_one({"id": box["id"]}, {"_id": 0})


@router.delete("/cashbox/boxes/{box_id}")
async def delete_box(box_id: str, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    if await db.cash_box_movements.count_documents({"box_id": box["id"]}) > 0:
        raise HTTPException(
            status_code=409,
            detail="La caja tiene movimientos; desactívala en lugar de borrarla")
    await db.cash_boxes.delete_one({"id": box["id"]})
    await db.cash_box_arqueos.delete_many({"box_id": box["id"]})
    return {"ok": True}


# ------------------------------------------------------- saldo inicial

@router.post("/cashbox/boxes/{box_id}/inicial")
async def set_initial(box_id: str, payload: InitialSet, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    denoms = _clean_denoms(payload.fund, payload.denominations, payload.amount)
    # H07 — control contable: con operativa registrada el inicial es inmutable;
    # las correcciones van como movimientos fechados de entrada/salida.
    has_movs = await db.cash_box_movements.count_documents(
        {"box_id": box["id"], "fund": payload.fund})
    has_arqueos = await db.cash_box_arqueos.count_documents(
        {"box_id": box["id"], "fund": payload.fund})
    if has_movs or has_arqueos:
        raise HTTPException(
            status_code=409,
            detail=("El fondo ya tiene movimientos o arqueos: el saldo inicial "
                    "no puede reescribirse. Registra una entrada/salida de "
                    "corrección con su concepto."))
    prev = (box.get("initial") or {}).get(payload.fund) or {}
    ops: Dict[str, Any] = {"$set": {
        f"initial.{payload.fund}": {
            "amount": round(float(payload.amount), 2),
            "denominations": denoms,
            "set_at": iso(now_utc()),
            "set_by": user.get("name") or user.get("email") or "",
        }}}
    if prev.get("set_at"):
        # historial auditable del valor anterior (H07)
        ops["$push"] = {f"initial_history.{payload.fund}": {
            "prev_amount": float(prev.get("amount") or 0),
            "new_amount": round(float(payload.amount), 2),
            "at": iso(now_utc()),
            "by": user.get("name") or user.get("email") or "",
        }}
    await db.cash_boxes.update_one({"id": box["id"]}, ops)
    return await db.cash_boxes.find_one({"id": box["id"]}, {"_id": 0})


# --------------------------------------------------------- movimientos

@router.get("/cashbox/boxes/{box_id}/movimientos")
async def list_movements(box_id: str, fund: str, request: Request,
                         month: Optional[str] = None) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    if fund not in FUNDS:
        raise HTTPException(status_code=400, detail="Fondo inválido")
    q: Dict[str, Any] = {"box_id": box["id"], "fund": fund}
    if month:
        start, end = _month_range_utc(month)
        q["created_at"] = {"$gte": start, "$lt": end}
    return await db.cash_box_movements.find(q, {"_id": 0}) \
        .sort("created_at", -1).to_list(2000)


@router.post("/cashbox/boxes/{box_id}/movimientos")
async def create_movement(box_id: str, payload: MovementCreate,
                          request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    denoms = _clean_denoms(payload.fund, payload.denominations, payload.amount)
    doc = {
        "id": f"cmov_{uuid.uuid4().hex[:12]}",
        "box_id": box["id"],
        "fund": payload.fund,
        "type": payload.type,
        "amount": round(float(payload.amount), 2),
        "concept": payload.concept.strip(),
        "responsible": payload.responsible.strip(),
        "denominations": denoms,
        "created_at": iso(now_utc()),
        "created_by_id": user["user_id"],
        "created_by_name": user.get("name") or user.get("email") or "",
    }
    # H01 — movimiento directo en la caja del sistema: queda marcado como
    # diferencia sin contrapartida contable, pendiente de conciliación.
    if box.get("system_purpose") == "company_cash":
        doc["ledger_status"] = "sin_contrapartida"
    await db.cash_box_movements.insert_one({**doc})
    await _bump_fund_rev(box["id"], payload.fund)
    return doc


@router.put("/cashbox/boxes/{box_id}/movimientos/{mov_id}")
async def update_movement(box_id: str, mov_id: str, payload: MovementUpdate,
                          request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    mov = await db.cash_box_movements.find_one(
        {"id": mov_id, "box_id": box["id"]}, {"_id": 0})
    if not mov:
        raise HTTPException(status_code=404, detail="Movimiento no encontrado")
    if _ledger_linked(mov):
        # H01 — espejo del Fondo de Empresa: importe/concepto vienen del
        # documento origen; solo se permite COMPLETAR el desglose pendiente.
        if (payload.denominations is None or payload.amount is not None
                or payload.concept is not None
                or payload.responsible is not None
                or mov.get("denominations")):
            raise HTTPException(
                status_code=409,
                detail=("Este movimiento proviene del Fondo de Empresa: solo "
                        "puede completarse su desglose de billetes pendiente."))
        denoms = _clean_denoms(mov["fund"], payload.denominations,
                               mov["amount"])
        await db.cash_box_movements.update_one(
            {"id": mov_id},
            {"$set": {"denominations": denoms, "denoms_pending": False}})
        await _bump_fund_rev(box["id"], mov["fund"])
        return await db.cash_box_movements.find_one({"id": mov_id}, {"_id": 0})
    upd: Dict[str, Any] = {}
    if payload.amount is not None:
        upd["amount"] = round(float(payload.amount), 2)
    if payload.concept is not None:
        upd["concept"] = payload.concept.strip()
    if payload.responsible is not None:
        upd["responsible"] = payload.responsible.strip()
    if payload.denominations is not None:
        upd["denominations"] = _clean_denoms(
            mov["fund"], payload.denominations,
            upd.get("amount", mov["amount"]))
    elif payload.amount is not None and mov.get("denominations"):
        # importe cambiado sin nuevo desglose → el viejo ya no cuadra
        upd["denominations"] = None
    if upd:
        await db.cash_box_movements.update_one({"id": mov_id}, {"$set": upd})
        await _bump_fund_rev(box["id"], mov["fund"])
    return await db.cash_box_movements.find_one({"id": mov_id}, {"_id": 0})


@router.delete("/cashbox/boxes/{box_id}/movimientos/{mov_id}")
async def delete_movement(box_id: str, mov_id: str, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    mov = await db.cash_box_movements.find_one(
        {"id": mov_id, "box_id": box["id"]}, {"_id": 0})
    if not mov:
        raise HTTPException(status_code=404, detail="Movimiento no encontrado")
    if _ledger_linked(mov):
        raise HTTPException(
            status_code=409,
            detail=("Este movimiento proviene del Fondo de Empresa y no puede "
                    "eliminarse desde la caja."))
    await db.cash_box_movements.delete_one({"id": mov_id, "box_id": box["id"]})
    await _bump_fund_rev(box["id"], mov["fund"])
    return {"ok": True}


# ------------------------------------------------------------- resumen

@router.get("/cashbox/boxes/{box_id}/resumen")
async def fund_summary(box_id: str, fund: str, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    if fund not in FUNDS:
        raise HTTPException(status_code=400, detail="Fondo inválido")
    # iter254(R11) — totales y denominaciones por AGREGACIÓN (sin cargar los
    # movimientos en memoria ni truncar en 20k: el resumen nunca subestima).
    base = {"box_id": box["id"], "fund": fund}
    initial = (box.get("initial") or {}).get(fund) or {}
    init_amount = float(initial.get("amount") or 0)

    tot = {r["_id"]: float(r["total"]) async for r in
           db.cash_box_movements.aggregate([
               {"$match": base},
               {"$group": {"_id": "$type", "total": {"$sum": "$amount"}}}])}
    entradas = round(tot.get("entrada", 0.0), 2)
    salidas = round(tot.get("salida", 0.0), 2)

    month = datetime.now(_TZ).strftime("%Y-%m")
    start, end = _month_range_utc(month)
    mtot = {r["_id"]: float(r["total"]) async for r in
            db.cash_box_movements.aggregate([
                {"$match": {**base, "created_at": {"$gte": start, "$lt": end}}},
                {"$group": {"_id": "$type", "total": {"$sum": "$amount"}}}])}
    e_mes = round(mtot.get("entrada", 0.0), 2)
    s_mes = round(mtot.get("salida", 0.0), 2)

    # control físico teórico por denominación (inicial ± desgloses)
    counts: Dict[str, int] = {}
    for d, q in (initial.get("denominations") or {}).items():
        counts[d] = counts.get(d, 0) + int(q)
    async for row in db.cash_box_movements.aggregate([
        {"$match": {**base, "denominations": {"$type": "object"}}},
        {"$project": {"type": 1,
                      "den": {"$objectToArray": "$denominations"}}},
        {"$unwind": "$den"},
        {"$group": {"_id": "$den.k", "qty": {"$sum": {
            "$cond": [{"$eq": ["$type", "entrada"]},
                      "$den.v", {"$multiply": ["$den.v", -1]}]}}}},
    ]):
        d = str(row["_id"])
        counts[d] = counts.get(d, 0) + int(row["qty"] or 0)
    denom_rows = [{"denom": d, "qty": counts.get(str(d), 0),
                   "subtotal": round(d * counts.get(str(d), 0), 2)}
                  for d in DENOMS[fund]]
    num_movs = await db.cash_box_movements.count_documents(base)

    # iter264/H04 — el arqueo solo vale como cierre si no hubo movimientos después
    from services.cash_box_arqueo import (closing_arqueo_status,
                                          havana_day_start_utc)
    day_start = havana_day_start_utc()
    last_arq, vigente = await closing_arqueo_status(box["id"], fund)
    arqueo_today = {**last_arq, "superseded": not vigente} if last_arq else None
    movs_today = await db.cash_box_movements.count_documents(
        {**base, "created_at": {"$gte": day_start}})
    balance = round(init_amount + entradas - salidas, 2)

    # H01 — movimientos directos sin contrapartida contable (a conciliar)
    sc: Dict[str, Any] = {"count": 0, "entradas": 0.0, "salidas": 0.0}
    async for row in db.cash_box_movements.aggregate([
        {"$match": {**base, "ledger_status": "sin_contrapartida"}},
        {"$group": {"_id": "$type", "total": {"$sum": "$amount"},
                    "n": {"$sum": 1}}},
    ]):
        sc["count"] += int(row.get("n") or 0)
        key = "entradas" if row["_id"] == "entrada" else "salidas"
        sc[key] = round(float(row.get("total") or 0), 2)

    return {
        "fund": fund,
        "initial": initial,
        "balance": balance,
        "entradas_total": entradas,
        "salidas_total": salidas,
        "month": month,
        "entradas_mes": e_mes,
        "salidas_mes": s_mes,
        "neto_mes": round(e_mes - s_mes, 2),
        "denominaciones": denom_rows,
        "num_movimientos": num_movs,
        "arqueo_today": arqueo_today,
        "needs_arqueo": (not vigente) and bool(
            movs_today or abs(balance) > 0.009),
        "sin_contrapartida": sc,
    }


# -------------------------------------------------------------- arqueo

@router.get("/cashbox/boxes/{box_id}/arqueos")
async def list_arqueos(box_id: str, fund: str, request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    return await db.cash_box_arqueos.find(
        {"box_id": box["id"], "fund": fund}, {"_id": 0}) \
        .sort("created_at", -1).to_list(100)


@router.post("/cashbox/boxes/{box_id}/arqueos")
async def create_arqueo(box_id: str, payload: ArqueoCreate,
                        request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    counted = _clean_denoms(payload.fund, payload.counted) or {}
    total = round(sum(int(float(d)) * q for d, q in counted.items()), 2)
    system = await _fund_balance(box, payload.fund)
    diff = round(total - system, 2)
    doc = {
        "id": f"carq_{uuid.uuid4().hex[:12]}",
        "box_id": box["id"],
        "fund": payload.fund,
        "counted": counted,
        "total_counted": total,
        "system_balance": system,
        "difference": diff,
        "status": "cuadrada" if diff == 0 else ("sobrante" if diff > 0 else "faltante"),
        # V04 — revisión del fondo al momento del conteo: cualquier mutación
        # posterior de movimientos deja este cierre como histórico
        "fund_rev": int(((box.get("fund_revs") or {}).get(payload.fund)) or 0),
        "note": payload.note.strip(),
        "created_at": iso(now_utc()),
        "created_by_id": user["user_id"],
        "created_by_name": user.get("name") or user.get("email") or "",
    }
    await db.cash_box_arqueos.insert_one({**doc})
    # iter264 — faltante/sobrante en caja de empresa → alerta a los admins
    if doc["status"] != "cuadrada":
        from services.cash_box_arqueo import notify_arqueo_discrepancy
        await notify_arqueo_discrepancy(box, doc)
    return doc


# ------------------------------------------------------------- reporte

@router.get("/cashbox/boxes/{box_id}/reporte")
async def monthly_report(box_id: str, fund: str, month: str,
                         request: Request) -> Any:
    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    if fund not in FUNDS:
        raise HTTPException(status_code=400, detail="Fondo inválido")
    try:
        datetime.strptime(month + "-01", "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail="Mes inválido (YYYY-MM)")
    start, end = _month_range_utc(month)
    days: Dict[str, Dict[str, float]] = {}
    num_movs = 0
    # H06 — iteración por cursor: sin tope silencioso de 20.000 movimientos
    async for m in db.cash_box_movements.find(
            {"box_id": box["id"], "fund": fund,
             "created_at": {"$gte": start, "$lt": end}},
            {"_id": 0, "type": 1, "amount": 1, "created_at": 1}):
        num_movs += 1
        local_day = datetime.fromisoformat(m["created_at"]) \
            .astimezone(_TZ).strftime("%Y-%m-%d")
        row = days.setdefault(local_day, {"entradas": 0.0, "salidas": 0.0})
        key = "entradas" if m["type"] == "entrada" else "salidas"
        row[key] += m["amount"]
    rows: list = [{"date": d,
                   "entradas": round(v["entradas"], 2),
                   "salidas": round(v["salidas"], 2),
                   "neto": round(v["entradas"] - v["salidas"], 2)}
                  for d, v in sorted(days.items())]
    return {
        "month": month,
        "days": rows,
        "total_entradas": round(sum(r["entradas"] for r in rows), 2),
        "total_salidas": round(sum(r["salidas"] for r in rows), 2),
        "neto": round(sum(r["neto"] for r in rows), 2),
        "num_movimientos": num_movs,
    }


# -------------------------------------------------------------- export

@router.get("/cashbox/boxes/{box_id}/export.xlsx")
async def export_xlsx(box_id: str, fund: str, request: Request) -> Any:
    from io import BytesIO
    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Font

    user = await require_user(request)
    box = await _get_box_checked(box_id, user)
    if fund not in FUNDS:
        raise HTTPException(status_code=400, detail="Fondo inválido")
    initial = (box.get("initial") or {}).get(fund) or {}

    wb = Workbook()
    ws = wb.active
    ws.title = f"Caja {fund}"
    bold = Font(bold=True)
    ws.append([f"Caja: {box['name']} — Fondo {fund}"])
    ws["A1"].font = bold
    ws.append([f"Saldo inicial: {float(initial.get('amount') or 0):g}"])
    ws.append([])
    ws.append(["Fecha", "Tipo", "Importe", "Concepto", "Responsable",
               "Registrado por", "Saldo acumulado"])
    for c in ws[4]:
        c.font = bold
    running = float(initial.get("amount") or 0)
    # H06 — iteración por cursor: el respaldo incluye TODOS los movimientos
    async for m in db.cash_box_movements.find(
            {"box_id": box["id"], "fund": fund}, {"_id": 0}) \
            .sort("created_at", 1):
        running += m["amount"] if m["type"] == "entrada" else -m["amount"]
        local = datetime.fromisoformat(m["created_at"]).astimezone(_TZ)
        ws.append([local.strftime("%Y-%m-%d %H:%M"),
                   m["type"].capitalize(), m["amount"], m["concept"],
                   m.get("responsible") or "", m.get("created_by_name") or "",
                   round(running, 2)])
    for col, w in zip("ABCDEFG", (17, 10, 12, 34, 18, 18, 15)):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"caja_{fund}_{datetime.now(_TZ).strftime('%Y-%m-%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
